import io
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from server import app as api
from server.word_export import compose_docx
from server.tests.test_letterheads import paper


class WordExportTests(unittest.TestCase):
    def setUp(self):
        self.temp=TemporaryDirectory();self.paths=api.DATA_DIR,api.UPLOAD_DIR,api.DB_PATH
        api.DATA_DIR=Path(self.temp.name);api.UPLOAD_DIR=api.DATA_DIR/'uploads';api.DB_PATH=api.DATA_DIR/'app.db';api.failed_logins.clear()

    def tearDown(self):
        api.DATA_DIR,api.UPLOAD_DIR,api.DB_PATH=self.paths;self.temp.cleanup()

    def login(self,client,email='word-test@example.invalid'):
        client.headers['X-Schimmelpilz-Request']='1'
        response=client.post('/api/auth/register',json=dict(email=email,password='Synthetic-word-test-2026!',audience='consumer',full_name='Synthetic Person',street='Testweg 1',postal_code='10115',city='Berlin',country='Deutschland'))
        self.assertEqual(response.status_code,201,response.text)

    def upload(self,client,label):
        head=client.post('/api/auth/letterheads',files={'file':('paper.pdf',paper(label),'application/pdf')}).json()
        self.assertEqual(client.put(f'/api/auth/letterheads/{head["id"]}/activate',json={'confirmed':True}).status_code,200)
        return head

    def test_editable_text_geometry_and_two_backgrounds(self):
        path=Path(self.temp.name)/'paper.pdf';path.write_bytes(paper())
        body='Betreff Test\n\nÄÖÜ ß € <literal>\n'+ '\n'.join('Unique line '+str(i) for i in range(90))
        content=compose_docx(body,path,api.LetterheadLayout().model_dump(),'TEST-1',2)
        doc=Document(io.BytesIO(content))
        self.assertEqual('\n'.join(p.text for p in doc.paragraphs),body)
        self.assertAlmostEqual(doc.sections[0].page_width.mm,210,places=1)
        self.assertAlmostEqual(doc.sections[0].top_margin.mm,70,places=1)
        self.assertTrue(doc.sections[0].different_first_page_header_footer)
        self.assertNotIn('approved',doc.core_properties.keywords)
        with ZipFile(io.BytesIO(content)) as archive:
            self.assertEqual(len([name for name in archive.namelist() if name.startswith('word/media/')]),2)
            headers=[archive.read(name).decode() for name in archive.namelist() if name.startswith('word/header') and name.endswith('.xml')]
            self.assertEqual(len(headers),2)
            self.assertTrue(all('behindDoc="1"' in value and 'relativeFrom="page"' in value and 'wrapNone' in value for value in headers))
            self.assertFalse(any(b'TargetMode="External"' in archive.read(name) for name in archive.namelist() if name.endswith('.rels')))

    def test_plain_docx_and_invalid_frame(self):
        content=compose_docx('Plain editable text',None,api.LetterheadLayout(top_mm=25,bottom_mm=25).model_dump())
        self.assertEqual(Document(io.BytesIO(content)).paragraphs[0].text,'Plain editable text')
        with self.assertRaises(ValueError):
            compose_docx('Test',None,api.LetterheadLayout(top_mm=160,bottom_mm=120).model_dump())
        with self.assertRaises(ValueError):
            compose_docx('Bad\x00text',None,api.LetterheadLayout().model_dump())

    def test_saved_snapshot_unsaved_editor_and_account_isolation(self):
        with TestClient(api.app) as client:
            self.login(client)
            first=self.upload(client,'OLD HEADER')
            case=client.post('/api/cases',json={'title':'Synthetic Word case','description':'Test only.'}).json()
            draft=client.post(f'/api/cases/{case["id"]}/drafts',json={'body':'Saved body ÄÖÜ ß €','kind':'letter'}).json()
            endpoint=f'/api/cases/{case["id"]}/drafts/{draft["id"]}/docx'
            original=client.get(endpoint)
            self.assertEqual(original.status_code,200,original.text if original.status_code!=200 else '')
            self.assertEqual(original.headers['X-Draft-Version'],'1')
            self.assertEqual(original.headers['X-Draft-Review-Status'],'editable-export')
            self.upload(client,'NEW HEADER')
            retained=client.get(endpoint)
            with ZipFile(io.BytesIO(original.content)) as old,ZipFile(io.BytesIO(retained.content)) as new:
                self.assertEqual(old.read('word/media/image1.png'),new.read('word/media/image1.png'))
            preview=client.post(f'/api/cases/{case["id"]}/drafts/export-docx',json={'draft_id':draft['id'],'body':'Unsaved editor text','use_current_letterhead':True})
            self.assertEqual(preview.status_code,200,preview.text if preview.status_code!=200 else '')
            self.assertEqual(Document(io.BytesIO(preview.content)).paragraphs[0].text,'Unsaved editor text')
            with ZipFile(io.BytesIO(preview.content)) as new,ZipFile(io.BytesIO(original.content)) as old:
                self.assertNotEqual(new.read('word/media/image1.png'),old.read('word/media/image1.png'))
            self.assertEqual(client.get('/api/cases/'+case['id']).json()['drafts'][0]['body'],'Saved body ÄÖÜ ß €')
            self.login(client,'other-word@example.invalid')
            self.assertEqual(client.get(endpoint).status_code,404)
            self.assertEqual(client.post(f'/api/cases/{case["id"]}/drafts/export-docx',json={'body':'Foreign'}).status_code,404)


if __name__=='__main__':
    unittest.main()
