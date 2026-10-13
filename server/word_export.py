"""Editable OOXML body with a private rasterized letter-paper background."""
import io
import threading

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt
import pypdfium2 as pdfium

# PDFium forbids concurrent calls, even on different documents.
PDFIUM_LOCK = threading.Lock()


def page_images(template):
    with PDFIUM_LOCK:
        with pdfium.PdfDocument(str(template)) as pdf:
            if not 1 <= len(pdf) <= 2:
                raise ValueError('invalid_letterhead')
            width, height = pdf[0].get_size()
            if width * height * 2.5 ** 2 > 10000000:
                raise ValueError('invalid_letterhead')
            images = []
            for index in range(len(pdf)):
                page = pdf[index]
                try:
                    bitmap = page.render(scale=2.5)
                    try:
                        image = bitmap.to_pil()
                        output = io.BytesIO()
                        image.save(output, format='PNG')
                        image.close()
                        images.append(output.getvalue())
                    finally:
                        bitmap.close()
                finally:
                    page.close()
            return images, width * 25.4 / 72, height * 25.4 / 72


def background(header, content, width, height):
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    run = paragraph.add_run()
    run.font.size = Pt(1)
    inline = run.add_picture(io.BytesIO(content), width=Mm(width), height=Mm(height))._inline
    anchor = OxmlElement('wp:anchor')
    for key, value in {'distT':'0', 'distB':'0', 'distL':'0', 'distR':'0', 'simplePos':'0', 'relativeHeight':'0', 'behindDoc':'1', 'locked':'1', 'layoutInCell':'1', 'allowOverlap':'1'}.items():
        anchor.set(key, value)
    simple = OxmlElement('wp:simplePos');simple.set('x','0');simple.set('y','0');anchor.append(simple)
    for axis in ['H','V']:
        position = OxmlElement('wp:position'+axis);position.set('relativeFrom','page')
        offset = OxmlElement('wp:posOffset');offset.text='0';position.append(offset);anchor.append(position)
    anchor.append(inline.extent)
    effect = OxmlElement('wp:effectExtent')
    for key in ['l','t','r','b']:
        effect.set(key,'0')
    anchor.append(effect);anchor.append(OxmlElement('wp:wrapNone'))
    inline.docPr.set('name','Letter paper background')
    inline.docPr.set('descr','Fixed image of the saved letter paper; document body remains editable')
    for child in list(inline):
        anchor.append(child)
    inline.getparent().replace(inline, anchor)


def compose_docx(body, template, layout, reference='', version=None):
    images, width, height = page_images(template) if template else ([], 210, 297)
    left,right,top,bottom = (layout[name] for name in ['left_mm','right_mm','top_mm','bottom_mm'])
    if width-left-right < 70 or height-top-bottom < 70:
        raise ValueError('body_frame_too_small')
    doc = Document()
    section = doc.sections[0]
    section.page_width,section.page_height = Mm(width),Mm(height)
    section.left_margin,section.right_margin,section.top_margin,section.bottom_margin = Mm(left),Mm(right),Mm(top),Mm(bottom)
    section.header_distance=Mm(0);section.footer_distance=Mm(0)
    if images:
        section.different_first_page_header_footer = len(images)==2
        if len(images)==2:
            background(section.first_page_header,images[0],width,height)
        background(section.header,images[-1],width,height)
    normal=doc.styles['Normal']
    normal.font.name='Arial';normal.font.size=Pt(layout.get('font_size',11))
    normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.line_spacing=1.45
    normal.paragraph_format.widow_control=False
    for line in body.replace('\r\n','\n').replace('\r','\n').split('\n'):
        doc.add_paragraph(line)
    doc.core_properties.title='Schimmelpilz '+reference
    doc.core_properties.author=''
    doc.core_properties.last_modified_by=''
    doc.core_properties.comments='Editable draft export. No approval for changes made in Word. Letter paper is a raster image; use PDF for the original preserved background.'
    doc.core_properties.keywords=f'draft; editable; version={version if version is not None else "unsaved"}'
    output=io.BytesIO();doc.save(output)
    return output.getvalue()
